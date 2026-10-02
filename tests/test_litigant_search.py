"""Attorney/litigant routing and exports with synthetic official-API-shaped replies."""
import contextlib
import csv
import io
import tempfile
from pathlib import Path
import unittest
from unittest.mock import patch
import zipfile

from openpyxl import load_workbook
from roc.cli import run
from roc.common import RocError, write_json
from roc.pacer import Session
from roc.search import search_criteria, subject_name
from roc.workspace import Workspace, search_config
from tests.test_workspace import record, response, form


def litigant_record(last='Acme Corporation', first='', seq=1, role='dft'):
    case = record(seq=seq, kind='Civil')
    case.update(caseTitle='Example v. ' + last, natureOfSuit='190')
    return {'lastName':last, 'firstName':first, 'middleName':'', 'partyType':'pty', 'partyRole':role, 'courtCase':case}


class LitigantSearchTests(unittest.TestCase):
    def test_unexpected_attorney_record_stops_without_hiding_receipts_or_error(self):
        calls = []
        def request(url, payload, headers):
            calls.append(payload)
            return response([litigant_record() | {'partyType':'aty'}])
        with tempfile.TemporaryDirectory() as folder, contextlib.redirect_stdout(io.StringIO()):
            ws = Workspace(folder, lambda:Session('fake-token', requester=request))
            try:
                ident = ws.new({'searchType':'litigant','lastName':'Acme Corporation','budgetCents':100})
                ws.future.result(15)
                summary = ws.summary(ident, True)
                self.assertEqual(summary['state'], 'stopped')
                self.assertIn('attorney record appeared', summary['message'])
                self.assertEqual(summary['spentCents'], 10)
                self.assertEqual(summary['cases'], [])
                self.assertEqual(summary['downloads'], [])
                self.assertEqual(calls, [{'lastName':'Acme Corporation','partyType':'pty'}])
            finally:
                ws.close()

    def test_explicit_party_type_and_optional_first_name_and_legacy_attorney_default(self):
        entity = search_config({'searchType':'litigant', 'lastName':' Acme Corporation ', 'budgetCents':100,
                                'courts':['nysdc'], 'dateFiledFrom':'2020-01-01'})
        self.assertNotIn('lawyer', entity)
        self.assertEqual(subject_name(entity), 'Acme Corporation')
        self.assertEqual(search_criteria(entity), {'lastName':'Acme Corporation', 'partyType':'pty',
                         'courtCase':{'courtId':['nysdc'], 'dateFiledFrom':'2020-01-01'}})
        person = search_config(form() | {'searchType':'litigant'})
        self.assertEqual(search_criteria(person), {'lastName':'Lawyer', 'firstName':'Jordan', 'partyType':'pty'})
        self.assertEqual(search_criteria({'lawyer':form()}), {'lastName':'Lawyer', 'firstName':'Jordan', 'partyType':'aty'})
        for invalid in (form() | {'searchType':'unknown'}, form() | {'firstName':''},
                        form() | {'searchType':'litigant','lastName':''},
                        form() | {'searchType':'litigant','aliases':['Other Name']},
                        form() | {'searchType':'litigant','partyType':'aty'}):
            with self.assertRaises(RocError):
                search_config(invalid)

    def test_fresh_entity_search_deduplicates_cases_preserves_names_roles_and_exports(self):
        calls = []
        rows = [litigant_record(), litigant_record(role='pla'), litigant_record(last='Acme Corporation II', seq=2, role='custom-code')]
        def request(url, payload, headers):
            calls.append((url, payload))
            return response(rows)
        with tempfile.TemporaryDirectory() as folder, contextlib.redirect_stdout(io.StringIO()), \
                patch('roc.cli.CourtRetriever', side_effect=AssertionError('No court retrieval')), \
                patch('roc.claude.analyze', side_effect=AssertionError('No AI')):
            ws = Workspace(folder, lambda:Session('fake-token', requester=request))
            try:
                ident = ws.new({'searchType':'litigant','lastName':'Acme Corporation','budgetCents':100})
                ws.future.result(15)
                summary = ws.summary(ident, True)
                self.assertEqual(summary['state'], 'ready')
                self.assertEqual(summary['searchType'], 'litigant')
                self.assertEqual(summary['subjectName'], 'Acme Corporation')
                self.assertNotIn('lawyer', summary)
                self.assertEqual(summary['caseCount'], 2)
                self.assertEqual(summary['spentCents'], 10)
                self.assertEqual(len(calls), 1)
                self.assertTrue(calls[0][0].startswith('https://pcl.uscourts.gov/pcl-public-api/rest/parties/find'))
                self.assertEqual(calls[0][1], {'lastName':'Acme Corporation', 'partyType':'pty'})
                self.assertEqual({p['role'] for c in summary['cases'] for p in c['indexedParties']}, {'pla','dft','custom-code'})
                self.assertFalse(any(c['eligible'] for c in summary['cases']))
                self.assertEqual(summary['partyReports']['clients']['summary'], [])
                output = ws.folder(ident) / 'output'
                with (output / 'case-index.csv').open(encoding='utf-8-sig', newline='') as stream:
                    exported = list(csv.reader(stream))
                self.assertEqual(len(exported), 3)
                book = load_workbook(output / 'case-index.xlsx')
                try:
                    self.assertIn('Matched litigants', book.sheetnames)
                    self.assertNotIn('Clients', book.sheetnames)
                    self.assertNotIn('Clients cases', book.sheetnames)
                    self.assertEqual({row[0].value for row in book['Matched litigants'].iter_rows(min_row=4)}, {'Acme Corporation','Acme Corporation II'})
                    self.assertEqual(book['Matched litigants'].max_row, 6)
                finally:
                    book.close()
                with zipfile.ZipFile(output / 'party-reports.zip') as bundle:
                    self.assertIn('matched-litigants.csv', bundle.namelist())
                    self.assertNotIn('clients-summary.csv', bundle.namelist())
                for fn in (lambda:ws.quote(ident, [summary['cases'][0]['key']]),
                           lambda:ws.grabber.analysis_quote(ident, {'keys':[summary['cases'][0]['key']]}),
                           lambda:ws.grabber.state(ident)):
                    with self.assertRaisesRegex(RocError, 'not available for litigants'):
                        fn()
                # Saved runs can reopen and rebuild without any new authentication/search.
                ws.close()
                ws = Workspace(folder, lambda:(_ for _ in ()).throw(AssertionError('No new sign-in')))
                self.assertEqual(ws.list()['jobs'][0]['searchType'], 'litigant')
                for action in ('export','refresh-reports'):
                    ws.act(ident, action)
                    ws.future.result(15)
                    self.assertEqual(ws.summary(ident)['state'], 'ready')
                self.assertEqual(len(calls), 1)
            finally:
                ws.close()

    def test_litigant_dockets_rejected_before_authentication_or_spending(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'config.json'
            for extra in ({'savedDockets':[{}]}, {'dockets':{'limit':1}},
                          {'retrieveDockets':[{'courtId':'nysdc','caseNumber':'1:24-cv-00001'}]}):
                write_json(path, {'searchType':'litigant','litigant':{'lastName':'Example Entity'},
                                 'runDirectory':'run','budgetCents':100} | extra)
                with patch.object(Session,'prompt', side_effect=AssertionError('No authentication')):
                    with self.assertRaisesRegex(RocError, 'not available for litigants'):
                        run(path, live=True)
            self.assertFalse((Path(folder) / 'run/ledger.json').exists())

    def test_litigant_pause_resume_reuses_identical_party_query_and_saved_page(self):
        calls = []
        with tempfile.TemporaryDirectory() as folder, contextlib.redirect_stdout(io.StringIO()), patch('roc.pacer.time.sleep'):
            def request(url, payload, headers):
                page = int(url.rsplit('=', 1)[1])
                calls.append((page, payload))
                if page == 0:
                    ws.pause_event.set()
                return response([litigant_record(last='Smith', first='Jane', seq=page+1)], page=page, last=page==1, total=2)
            ws = Workspace(folder, lambda:Session('fake-token', requester=request))
            try:
                ident = ws.new({'searchType':'litigant','firstName':'Jane','lastName':'Smith','budgetCents':100})
                ws.future.result(15)
                self.assertEqual(ws.summary(ident)['state'], 'stopped')
                self.assertEqual(ws.summary(ident)['spentCents'], 10)
                ws.act(ident, 'resume')
                ws.future.result(15)
                self.assertEqual(ws.summary(ident)['state'], 'ready')
                self.assertEqual([p for p,_ in calls], [0,1])
                self.assertTrue(all(c['partyType']=='pty' and c['firstName']=='Jane' for _,c in calls))
                self.assertEqual(ws.summary(ident)['caseCount'], 2)
            finally:
                ws.close()
