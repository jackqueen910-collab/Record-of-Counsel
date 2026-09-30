import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from roc.common import RocError, normalize_case_number
from roc.docket import enrich, parse_report, receipt_cents
from roc.index import build_index
from roc.output import csv_safe, export_local
from roc.pacer import Session, collect_index
from roc.retrieve import court_origin, select_case_line
from roc.store import RunStore


def party(role, name, lawyer, counts="", number="1"):
    role_text = role + (f" ({number})" if role == "Defendant" else "")
    return (f"<tr><td><b><u>{role_text}</u></b></td></tr>"
            f"<tr><td><b>{name}</b></td><td>represented by</td><td><b>{lawyer}</b></td></tr>"
            "<tr><td>Pending Counts</td><td></td><td>Disposition</td></tr>" + counts)


def count(label, ids, disposition=""):
    return f"<tr><td>{label}<br>({ids})</td><td></td><td>{disposition}</td></tr>"


def report(parties, kind="CRIMINAL", receipt=True):
    return (f"<h3>U.S. District Court<br>Example District<br>{kind} DOCKET FOR CASE #: 1:24-cr-00001-XYZ</h3><table>" + parties +
            "</table><table><tr><td>Date Filed</td><td>#</td><td>Docket Text</td></tr>"
            "<tr><td>01/01/2024</td><td>1</td><td>Another Attorney mentioned in text must not establish representation</td></tr></table>" +
            ("<table><tr><td>Transaction Receipt</td></tr><tr><td>Cost:</td><td>3.00</td></tr></table>" if receipt else ""))


class ParserTests(unittest.TestCase):
    def test_pro_se_party_is_not_an_attorney(self):
        html = report(party("Plaintiff", "Self Represented", "Self Represented<br>PRO SE"), "CIVIL")
        # In real CM/ECF, the marker follows the bold name, outside that tag.
        html = html.replace("<b>Self Represented<br>PRO SE</b>", "<b>Self Represented</b><br>PRO SE")
        parsed = parse_report(html)
        self.assertEqual(parsed["parties"][0]["counsel"], [])
        self.assertEqual(parsed["parties"][0]["selfRepresentedNames"], ["Self Represented"])
        self.assertEqual(enrich(parsed, ["Self Represented"])["team"], "")

    def test_pro_se_marker_does_not_remove_other_attorney_in_same_cell(self):
        html = report(party("Defendant", "Client", "Real Attorney", count("18:111.F ASSAULT", "1")))
        html = html.replace("<b>Real Attorney</b>", "<b>Client</b><br>PRO SE<br><b>Real Attorney</b><br>ATTORNEY TO BE NOTICED")
        parsed = parse_report(html)
        self.assertEqual(parsed["parties"][0]["counsel"], ["Real Attorney"])
        self.assertEqual(enrich(parsed, ["Real Attorney"])["team"], "Criminal Defense")

    def test_client_specific_not_all_defendants(self):
        html = report(party("Defendant", "Client", "Jordan Lawyer", count("18:1343.F FRAUD BY WIRE", "2s")) +
                      party("Defendant", "Other Person", "Other Attorney", count("18:111.F ASSAULT", "1s"), "2"))
        result = enrich(parse_report(html), ["Jordan Lawyer"])
        self.assertEqual(result["nature"], "Wire fraud (count 2).")
        self.assertEqual(result["representedParties"], ["Client"])
        self.assertEqual(result["team"], "Criminal Defense")

    def test_court_notification_contacts_are_not_counsel(self):
        html = report(party("Defendant", "Client", "Jordan Lawyer", count("18:111.F ASSAULT", "1")))
        html = html.replace("<b>Jordan Lawyer</b>",
            "<b>Jordan Lawyer</b><br>Office near Pretrial Services<br>Designation: Retained"
            "<b>Sample City Interpreter</b><br>Designation: Retained"
            "<b>Pretrial Office</b><br>ATTORNEY TO BE NOTICED<br>Designation: Pretrial Services"
            "<b>Probation Office</b><br>ATTORNEY TO BE NOTICED<br>Designation: Probation Department")
        parsed = parse_report(html)
        self.assertEqual(parsed["parties"][0]["counsel"], ["Jordan Lawyer"])
        self.assertEqual(parsed["parties"][0]["courtContacts"],
                         ["Sample City Interpreter", "Pretrial Office", "Probation Office"])
        self.assertEqual(enrich(parsed, ["Sample City Interpreter"])["team"], "")
        self.assertEqual(enrich(parsed, ["Jordan Lawyer"])["team"], "Criminal Defense")

    def test_superseded_versions_not_double_counted(self):
        html = report(party("Defendant", "Client", "Jordan Lawyer",
            count("18:1343.F FRAUD BY WIRE", "1", "Superseded") + count("18:1343.F FRAUD BY WIRE", "1s-3s", "Guilty")))
        result = enrich(parse_report(html), ["Jordan Lawyer"])
        self.assertEqual(result["nature"], "Wire fraud (counts 1–3).")
        self.assertNotIn("Guilty", result["nature"])
        self.assertNotIn("Client", result["nature"])

    def test_different_lawyer_does_not_match_initial(self):
        html = report(party("Defendant", "Client", "Jordan A. Lawyer", count("18:1343.F FRAUD BY WIRE", "1")))
        self.assertEqual(enrich(parse_report(html), ["Jordan B. Lawyer"])["team"], "")
        self.assertEqual(enrich(parse_report(html), ["Lawyer, Jordan A"])["team"], "Criminal Defense")

    def test_partial_supersession_requires_review(self):
        html = report(party("Defendant", "Client", "Jordan Lawyer", count("18:111.F ASSAULT", "1") + count("18:1343.F FRAUD BY WIRE", "2s")))
        result = enrich(parse_report(html), ["Jordan Lawyer"])
        self.assertEqual(result["nature"], "")
        self.assertTrue(result["warnings"])

    def test_granted_dismissal_motion_resolves_old_count_only_when_explicit(self):
        for disposition in ("govt's oral motion to dismiss granted", "Motion to dismiss was granted"):
            html = report(party("Defendant", "Client", "Jordan Lawyer",
                count("18:1343.F FRAUD BY WIRE", "1s") +
                count("18:111.F ASSAULT", "2", disposition)))
            result = enrich(parse_report(html), ["Jordan Lawyer"])
            self.assertEqual(result["nature"], "Wire fraud (count 1).")
            self.assertFalse(result["warnings"])
        for disposition in ("motion to dismiss denied", "motion to dismiss pending", "motion to dismiss not granted"):
            html = report(party("Defendant", "Client", "Jordan Lawyer",
                count("18:1343.F FRAUD BY WIRE", "1s") +
                count("18:111.F ASSAULT", "2", disposition)))
            result = enrich(parse_report(html), ["Jordan Lawyer"])
            self.assertEqual(result["nature"], "")
            self.assertTrue(result["warnings"])

    def test_multiple_clients_different_counts_require_review(self):
        html = report(party("Defendant", "A", "Jordan Lawyer", count("18:111.F ASSAULT", "1")) +
                      party("Defendant", "B", "Jordan Lawyer", count("18:1343.F FRAUD BY WIRE", "2"), "2"))
        result = enrich(parse_report(html), ["Jordan Lawyer"])
        self.assertEqual(result["nature"], "")
        self.assertEqual(len(result["selectedCounts"]), 2)

    def test_prosecution_uses_government_counsel(self):
        html = report(party("Defendant", "A", "Defense Counsel", count("18:111.F ASSAULT", "1")) + party("Plaintiff", "USA", "Jordan Lawyer"))
        result = enrich(parse_report(html), ["Jordan Lawyer"])
        self.assertEqual(result["team"], "Prosecution")
        self.assertEqual(result["representedParties"], ["USA"])
        self.assertEqual(result["nature"], "Assault (count 1).")

    def test_civil_sides(self):
        for role, expected in [("Defendant", "Civil Defense"), ("Plaintiff", "Civil Plaintiff")]:
            result = enrich(parse_report(report(party(role, "Client", "Jordan Lawyer"), "CIVIL")), ["Jordan Lawyer"])
            self.assertEqual(result["team"], expected)

    def test_absent_counts_not_guessed_from_docket_text(self):
        result = enrich(parse_report(report(party("Defendant", "Client", "Jordan Lawyer"))), ["Jordan Lawyer"])
        self.assertEqual(result["nature"], "")
        self.assertTrue(result["warnings"])

    def test_receipt_required(self):
        self.assertEqual(receipt_cents(report("")), 300)
        with self.assertRaises(RocError):
            receipt_cents(report("", receipt=False))

    def test_login_page_is_not_docket(self):
        with self.assertRaises(RocError):
            parse_report("<html>Sign in to PACER</html>")


class SpendingTests(unittest.TestCase):
    def test_budget_cache_and_duplicate_guard(self):
        with tempfile.TemporaryDirectory() as folder, RunStore(folder, 300) as store:
            t = store.reserve("docket", {"case": "1"}, 300)
            path = store.finish(t, report(""))
            self.assertEqual(store.cached("docket", {"case": "1"}), path)
            with self.assertRaises(RocError):
                store.reserve("docket", {"case": "2"}, 300)
            self.assertEqual(store.spent, 300)

    def test_timeout_blocks_retry_and_other_paid_requests(self):
        with tempfile.TemporaryDirectory() as folder:
            with RunStore(folder, 600) as store:
                store.reserve("docket", {"case": "1"}, 300)
            with RunStore(folder, 600) as store:
                with self.assertRaises(RocError):
                    store.cached("docket", {"case": "1"})
                with self.assertRaises(RocError):
                    store.reserve("docket", {"case": "2"}, 300)

    def test_saved_response_reconciles_without_network(self):
        with tempfile.TemporaryDirectory() as folder, RunStore(folder, 300) as store:
            t = store.reserve("docket", {"case": "1"}, 300)
            path = Path(folder) / t["responseFile"]
            path.parent.mkdir()
            path.write_text(report(""), encoding="utf-8")
            store.reconcile()
            self.assertEqual(store.spent, 300)

    def test_excess_receipt_stops_next_request(self):
        with tempfile.TemporaryDirectory() as folder, RunStore(folder, 1000) as store:
            t = store.reserve("docket", {"case": "1"}, 300)
            with self.assertRaises(RocError):
                store.finish(t, report("").replace("3.00", "3.10"))
            self.assertEqual(store.spent, 310)

    def test_concurrent_run_lock(self):
        with tempfile.TemporaryDirectory() as folder, RunStore(folder, 0):
            with self.assertRaises(RocError):
                with RunStore(folder, 0):
                    pass


class ApiTests(unittest.TestCase):
    def test_api_login_uses_otp_and_token_only_in_memory(self):
        def request(url, payload, headers=None):
            self.assertIn("/services/cso-auth", url)
            self.assertEqual(payload["otpCode"], "123456")
            return json.dumps({"loginResult": "0", "nextGenCSO": "test-token"}), {}
        self.assertEqual(Session.login("user", "secret", "123456", requester=request).token, "test-token")

    def test_failed_auth_is_not_retried(self):
        calls = []
        def request(*args):
            calls.append(1)
            return '{"loginResult":"13"}', {}
        with self.assertRaises(RocError):
            Session.login("user", "secret", requester=request)
        self.assertEqual(len(calls), 1)

    def test_pagination_receipts_and_cache(self):
        calls = []
        def request(url, payload, headers):
            n = int(url.split("page=")[1])
            calls.append(n)
            result = {"content": [{"page": n}], "pageInfo": {"number": n, "totalElements": 2, "numberOfElements": 1, "last": n == 1}, "receipt": {"searchFee": ".10"}}
            return json.dumps(result), {"X-NEXT-GEN-CSO": "rotated"}
        with tempfile.TemporaryDirectory() as folder, RunStore(folder, 20) as store:
            session = Session("token", requester=request)
            rows = collect_index(session, {"lastName": "Lawyer"}, store, delay=0)
            self.assertEqual(len(rows), 2)
            updates = []
            collect_index(session, {"lastName": "Lawyer"}, store, delay=0,
                          progress=lambda stage, message, **details: updates.append((message, details)))
            self.assertEqual(calls, [0, 1])
            self.assertEqual(session.token, "rotated")
            self.assertEqual(store.spent, 20)
            self.assertTrue(all("no new charge" in msg for msg, details in updates))
            self.assertTrue(all(details["newSearchChargesCents"] == 0 for msg, details in updates))


class IndexOutputTests(unittest.TestCase):
    def test_case_normalization(self):
        self.assertEqual(normalize_case_number("1:2018cr00217"), "1:18-cr-00217")
        self.assertEqual(normalize_case_number("1:18-cr-217-ABC-2"), "1:18-cr-00217")

    def test_court_is_part_of_dedup_key(self):
        rows = [{"courtId": court, "caseNumberFull": "1:2024cr00001", "caseTitle": "USA v. A", "jurisdictionType": "Criminal", "dateFiled": "2024-01-01"} for court in ("nysdc", "njdc")]
        self.assertEqual(len(build_index(rows)), 2)

    def test_main_case_selector_excludes_defendant_subcase(self):
        lines = ["1:24-cr-00001-XYZ-1 Client", "1:24-cr-00001-XYZ USA v. A (closed)"]
        self.assertEqual(select_case_line(lines, "1:24-cr-00001"), 1)
        with self.assertRaises(RocError):
            court_origin("https://ecf.nysd.uscourts.gov.example.org/cgi-bin/DktRpt.pl")

    def test_output_columns_and_formula_injection(self):
        from openpyxl import load_workbook
        case = build_index([{"courtId": "nysdc", "caseNumberFull": "1:2024cr00001", "caseTitle": "=HYPERLINK(\"bad\")", "jurisdictionType": "Criminal", "dateFiled": "2024-01-01"}])[0]
        with tempfile.TemporaryDirectory() as folder:
            path = export_local([case], folder, "Test", {"generatedUtc": "2026-09-29"})
            book = load_workbook(path)
            self.assertEqual(book.active.max_column, 10)
            self.assertEqual(book.active["B7"].data_type, "s")
            self.assertFalse(book.active["H7"].alignment.wrap_text)
            self.assertEqual(book.active["G7"].value.year, 2024)
            self.assertTrue(csv_safe("=1+1").startswith("'"))


if __name__ == "__main__":
    unittest.main()
