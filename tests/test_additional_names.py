"""Multiple explicit PCL names share one cache, ledger, and case index."""
import contextlib
import csv
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from openpyxl import load_workbook
from roc.exports import output_directory
from roc.common import RocError, read_json
from roc.pacer import Session
from roc.search import counsel_aliases, search_plan
from roc.workspace import Workspace, search_config
from tests.test_litigant_search import litigant_record
from tests.test_workspace import FakeCourt, form, record, response, retrieval_values


class AdditionalNamesTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.folder = Path(temp.name)
        self.calls = []
        self.addCleanup(patch.stopall)
        patch('roc.search_collection.time.sleep').start()
        self.stdout = contextlib.redirect_stdout(io.StringIO())
        self.stdout.__enter__()
        self.addCleanup(self.stdout.__exit__, None, None, None)

    def workspace(self, reply):
        def request(url, payload, headers):
            page = int(url.rsplit('=', 1)[1])
            self.calls.append((payload, page))
            return reply(payload, page)
        ws = Workspace(self.folder, lambda:Session('fake-token', requester=request))
        self.addCleanup(ws.close)
        return ws

    def test_validation_deduplication_and_legacy_docket_spellings(self):
        config = search_config(form() | {'aliases':['Jordan A. Lawyer'], 'additionalNames':[
            {'firstName':' Jordan ', 'lastName':'Lawyer'}, {'firstName':'Jordy', 'lastName':'Lawyer'},
            {'firstName':'Jordy', 'lastName':'Lawyer'}]})
        self.assertEqual([q['name'] for q in search_plan(config)], ['Jordan Lawyer','Jordy Lawyer'])
        self.assertEqual(counsel_aliases(config), ['Jordan Lawyer','Jordy Lawyer','Jordan A. Lawyer'])
        self.assertEqual(len(search_plan(search_config(form() | {'aliases':['Jordy Lawyer']}))), 1)
        for names in (None, 'Name', [{}], [{'lastName':'Lawyer'}], [{'lastName':'X', 'firstName':'Y','partyType':'pty'}], [{}]*31):
            with self.subTest(names=names), self.assertRaises(RocError):
                search_config(form() | {'additionalNames':names})
        entity = search_config({'searchType':'litigant','lastName':'Acme','budgetCents':100,
                               'additionalNames':[{'lastName':'The Acme Company'}, {'lastName':'Smith','firstName':'Jane'}]})
        self.assertNotIn('firstName', search_plan(entity)[1]['criteria'])
        self.assertEqual(search_plan(entity)[2]['criteria']['firstName'], 'Jane')

    def test_overlapping_entities_keep_one_case_and_both_query_sources_in_exports(self):
        rows = [litigant_record()]
        ws = self.workspace(lambda payload, page: response(rows + ([litigant_record(seq=2)] if payload['lastName']=='Acme Corporation' else [])))
        ident = ws.new({'searchType':'litigant','lastName':'Acme','budgetCents':100,
                        'courts':['nysdc'], 'dateFiledFrom':'2020-01-01',
                        'additionalNames':[{'lastName':'Acme Corporation'}, {'lastName':'Acme Corporation'}]})
        ws.future.result(15)
        summary = ws.summary(ident, True)
        self.assertEqual(summary['state'], 'ready')
        self.assertEqual(summary['caseCount'], 2)
        self.assertEqual(summary['spentCents'], 20)
        self.assertEqual(len(self.calls), 2)
        self.assertTrue(all(c[0]['partyType']=='pty' and c[0]['courtCase']=={'courtId':['nysdc'],'dateFiledFrom':'2020-01-01'} for c in self.calls))
        cases = ws.cases(ident)
        both = next(c for c in cases if c['caseNumber'].endswith('00001'))
        self.assertEqual(both['matchedSearchNames'], ['Acme','Acme Corporation'])
        self.assertEqual(len(both['sourceRows']), 1)
        self.assertEqual(both['indexedParties'][0]['name'], 'Acme Corporation')
        output = output_directory(ws.folder(ident))
        with zipfile.ZipFile(output/'party-reports.zip') as bundle:
            exported = list(csv.reader(io.StringIO(bundle.read('search-matches.csv').decode('utf-8-sig'))))
        self.assertEqual(len(exported), 4)
        book = load_workbook(output/'case-index.xlsx')
        self.assertIn('Search matches', book.sheetnames)
        book.close()
        for action in ('export','refresh-reports'):
            ws.act(ident, action)
            ws.future.result(15)
            self.assertEqual(ws.summary(ident)['state'], 'ready')
        self.assertEqual(len(self.calls), 2)
        self.assertEqual(ws.cases(ident)[-1]['matchedSearchNames'], ['Acme','Acme Corporation'])

    def test_shared_cap_stops_between_names_and_resume_never_rebuys(self):
        ws = self.workspace(lambda payload, page: response([litigant_record(last=payload['lastName'], seq=1 if payload['lastName']=='Acme' else 2)]))
        ident = ws.new({'searchType':'litigant','lastName':'Acme','budgetCents':10,
                        'additionalNames':[{'lastName':'Other Entity'}]})
        ws.future.result(15)
        summary = ws.summary(ident, True)
        self.assertEqual(summary['state'], 'stopped')
        self.assertIn('1/2 name searches', summary['message'])
        self.assertIn('Budget stop', summary['message'])
        self.assertFalse(summary['indexReady'])
        self.assertEqual(summary['spentCents'], 10)
        self.assertEqual(summary['caseCount'], 1)
        self.assertEqual(summary['downloads'], [])
        with self.assertRaises(RocError):
            ws.act(ident, 'export')
        ws.act(ident, 'resume')
        ws.future.result(15)
        self.assertEqual(len(self.calls), 1)
        ws.act(ident, 'budget', {'budgetCents':20})
        ws.act(ident, 'resume')
        ws.future.result(15)
        summary = ws.summary(ident, True)
        self.assertEqual((summary['state'], summary['caseCount'], summary['spentCents']), ('ready',2,20))
        self.assertEqual([c[0]['lastName'] for c in self.calls], ['Acme','Other Entity'])
        self.assertTrue(read_json(ws.folder(ident)/'search-progress.json')['complete'])

    def test_paginated_second_name_pause_and_resume_preserves_first_query(self):
        def reply(payload, page):
            second = payload['lastName']=='Other'
            if second and page==0:
                ws.pause_event.set()
            return response([litigant_record(last=payload['lastName'], seq=page+1)], page=page,
                            last=not second or page==1, total=2 if second else 1)
        ws = self.workspace(reply)
        ident = ws.new({'searchType':'litigant','lastName':'Acme','budgetCents':100,
                        'additionalNames':[{'lastName':'Other'}]})
        ws.future.result(15)
        self.assertEqual(ws.summary(ident)['state'], 'stopped')
        ws.act(ident, 'resume')
        ws.future.result(15)
        self.assertEqual(ws.summary(ident)['state'], 'ready')
        self.assertEqual([(c[0]['lastName'],c[1]) for c in self.calls], [('Acme',0),('Other',0),('Other',1)])
        self.assertEqual(ws.summary(ident)['spentCents'], 30)

    def test_uncertain_receipt_stops_remaining_names_and_retry(self):
        def reply(payload, page):
            if payload['lastName']=='Other':
                raise RocError('Request outcome unknown.')
            return response([litigant_record()])
        ws = self.workspace(reply)
        ident = ws.new({'searchType':'litigant','lastName':'Acme','budgetCents':100,
                        'additionalNames':[{'lastName':'Other'}, {'lastName':'Third'}]})
        ws.future.result(15)
        summary = ws.summary(ident, True)
        self.assertEqual((summary['state'], summary['pendingCount'], summary['spentCents']), ('stopped',1,10))
        ws.act(ident, 'resume')
        ws.future.result(15)
        self.assertEqual([c[0]['lastName'] for c in self.calls], ['Acme','Other'])

    def test_additional_attorney_name_flows_to_docket_clients_exports_and_document_sources(self):
        def reply(payload, page):
            return response([{'firstName':payload['firstName'],'lastName':'Lawyer','partyType':'aty','courtCase':record()}])
        ws = self.workspace(reply)
        ident = ws.new(form() | {'firstName':'Robert','additionalNames':[{'firstName':'Jordan','lastName':'Lawyer'}]})
        ws.future.result(15)
        summary = ws.summary(ident, True)
        self.assertEqual(summary['caseCount'], 1)
        self.assertEqual(summary['cases'][0]['matchedSearchNames'], ['Robert Lawyer','Jordan Lawyer'])
        key = summary['cases'][0]['key']
        FakeCourt.bought, FakeCourt.fail_key = [], None
        with patch('roc.engine.CourtRetriever', FakeCourt):
            ws.act(ident, 'retrieve', retrieval_values(ws, ident, [key]))
            ws.future.result(15)
        summary = ws.summary(ident, True)
        self.assertEqual(summary['cases'][0]['representedParties'], ['Example Client'])
        self.assertEqual(summary['partyReports']['clients']['summary'][0]['name'], 'Example Client')
        self.assertEqual(len(self.calls), 2)
        sources = ws.grabber.sources(ident, [key])
        self.assertTrue(sources[0]['clients'])
        self.assertEqual(read_json(output_directory(ws.folder(ident))/'party-reports.json')['clients']['summary'][0]['name'], 'Example Client')
        self.assertEqual(ws.cases(ident)[0]['matchedSearchNames'], ['Robert Lawyer','Jordan Lawyer'])
