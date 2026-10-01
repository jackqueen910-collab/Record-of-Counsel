"""Party/client association, distinct-case ranking and standalone output acceptance."""
import csv
import io
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from openpyxl import load_workbook

from roc.docket import enrich, parse_report
from roc.output import export_local, publish_google
from roc.parties import build_party_reports, party_key, report_tables
from tests.test_roc import party, count, report


def case(parties, key="nysdc|1:24-cv-00001", kind="CIVIL", aliases=("Jordan Lawyer",)):
    raw = report(parties, kind).replace("1:24-cr-00001", key.split("|")[1])
    if kind == "CIVIL":
        raw = raw.replace('</h3>', '</h3><p>Nature of Suit: 190 Contract Jurisdiction: Diversity</p>')
    parsed = parse_report(raw)
    evidence = enrich(parsed, aliases)
    evidence.update(parsedDocket=parsed, sourceFile="fixture.html")
    return {"key": key, "courtId": key.split("|")[0], "caseNumber": key.split("|")[1], "caseTitle": "Fictional case",
            "caseType": kind.title(), "court": "U.S. District Court", "district": "Example District", "dateFiled": "2024-01-01",
            "status": "Open", "pacerLink": "https://ecf.nysd.uscourts.gov/", "role": evidence["role"], "team": evidence["team"],
            "nature": evidence["nature"], "enrichment": evidence, "warnings": [], "issues": evidence["issues"]}


class PartyReportTests(unittest.TestCase):
    def test_multi_party_multi_counsel_repeats_and_distinct_court_cases(self):
        raw = (party("Plaintiff", "Ann Client", "Jordan Lawyer") + party("Plaintiff", "Second Client", "Jordan Lawyer") +
               party("Plaintiff", "Other Plaintiff", "Other Attorney") + party("Defendant", "ACME LLC", "Defense Lawyer") +
               party("Defendant", " acme  LLC ", "Defense Lawyer") + party("Defendant", "Acme Inc", "Defense Lawyer"))
        a = case(raw)
        b = case(raw, "njdc|1:24-cv-00001")  # Same number in another court is a separate case.
        c = case(party("Defendant", "ACME LLC", "Jordan Lawyer"), "nysdc|1:24-cv-00002")
        index_only = {k:v for k,v in a.items() if k != "enrichment"} | {"key": "nysdc|1:24-cv-00003"}
        reports = build_party_reports([a,b,c,index_only])
        self.assertEqual(reports["coverage"]["parsedDockets"], 3)
        self.assertEqual(reports["coverage"]["indexOnlyCases"], 1)
        s = reports["defendants"]["summary"][0]
        self.assertEqual(s["caseCount"], 3)
        self.assertEqual(s["opposingCaseCount"], 2)
        self.assertEqual(s["civilPlaintiffCaseCount"], 2)
        self.assertEqual(s["clientCaseCount"], 1)
        self.assertEqual(len(reports["defendants"]["cases"]), 5)
        self.assertEqual(len(s["sourceNames"]), 2)
        clients = {s["name"]:s["caseCount"] for s in reports["clients"]["summary"]}
        self.assertEqual(clients, {"Ann Client":2,"Second Client":2,"ACME LLC":1})
        other = next(s for s in reports["plaintiffs"]["summary"] if s["name"] == "Other Plaintiff")
        self.assertEqual(other["sameSideCaseCount"], 2)

    def test_criminal_counts_clients_and_codefendants_remain_separate(self):
        raw = (party("Plaintiff", "USA", "Government Lawyer") +
               party("Defendant", "Client", "Jordan Lawyer", count("CLIENT COUNT", "1s")) +
               party("Defendant", "Co-defendant", "Different Lawyer", count("OTHER COUNT", "2s"), "2"))
        defense = case(raw, kind="CRIMINAL")
        self.assertEqual(defense["nature"], "CLIENT COUNT (1s)")
        self.assertEqual(defense["role"], "Criminal Defense")
        r = build_party_reports([defense])
        other = next(s for s in r["defendants"]["summary"] if s["name"] == "Co-defendant")
        self.assertEqual(other["opposingCaseCount"], 0)
        self.assertEqual(other["sameSideCaseCount"], 1)
        self.assertEqual(r["plaintiffs"]["summary"][0]["opposingCaseCount"], 1)
        prosecution = case(raw, kind="CRIMINAL", aliases=["Government Lawyer"])
        r = build_party_reports([prosecution])
        self.assertEqual(r["clients"]["summary"][0]["name"], "USA")
        self.assertTrue(all(s["opposingCaseCount"] == 1 and s["civilPlaintiffCaseCount"] == 0 for s in r["defendants"]["summary"]))
        self.assertIn("OTHER COUNT", prosecution["nature"])

    def test_unknown_alias_keeps_parties_without_claiming_clients_or_opponents(self):
        c = case(party("Plaintiff", "Ann", "Another Lawyer") + party("Defendant", "Entity", "Defense"))
        r = build_party_reports([c])
        self.assertEqual(r["clients"]["summary"], [])
        self.assertEqual(r["defendants"]["summary"][0]["unresolvedCaseCount"], 1)
        self.assertEqual(r["coverage"]["casesWithoutMatchedClients"], 1)

    def test_defendant_numbers_preserve_namesake_client_boundaries(self):
        c = case(party("Defendant", "Same Name", "Jordan Lawyer", count("CLIENT COUNT", "1"), "1") +
                 party("Defendant", "Same Name", "Other Lawyer", count("OTHER COUNT", "2"), "2"), kind="CRIMINAL")
        r = build_party_reports([c])
        self.assertEqual(c["nature"], "CLIENT COUNT (1)")
        self.assertEqual(len(r["parties"]), 2)
        self.assertEqual(r["clients"]["cases"][0]["defendantNumbers"], ["1"])
        self.assertEqual(r["defendants"]["cases"][0]["defendantNumbers"], ["1", "2"])
        self.assertEqual(r["defendants"]["summary"][0]["caseCount"], 1)
        self.assertEqual(r["defendants"]["summary"][0]["sameSideCaseCount"], 1)
        self.assertEqual(r["defendants"]["summary"][0]["clientCaseCount"], 1)

    def test_legacy_evidence_requires_counsel_alias_not_same_party_name(self):
        c = case(party("Plaintiff", "Same Name", "Jordan Lawyer") + party("Defendant", "Same Name", "Other Lawyer"))
        del c["enrichment"]["partyDetails"]
        r = build_party_reports([c], ["Jordan Lawyer"])
        self.assertEqual(r["clients"]["cases"][0]["partyRole"], "Plaintiff")
        self.assertEqual(r["defendants"]["summary"][0]["clientCaseCount"], 0)
        self.assertEqual(build_party_reports([c])["clients"]["summary"], [])

    def test_partial_layout_flags_coverage_and_does_not_infer_opposition(self):
        c = case(party("Plaintiff", "Client", "Jordan Lawyer") + party("Defendant", "Entity", "Defense"))
        c["enrichment"]["parsedDocket"]["warnings"].append("Party table incomplete or unfamiliar.")
        r = build_party_reports([c])
        self.assertEqual(r["coverage"]["partyTablesNeedingReview"], 1)
        self.assertEqual(r["defendants"]["summary"][0]["unresolvedCaseCount"], 1)
        self.assertEqual(r["defendants"]["cases"][0]["partyTableStatus"], "Needs review")

    def test_mixed_roles_merge_client_case_rows_and_exclude_mediators(self):
        c = case(party("Petitioner", "Client", "Jordan Lawyer") + party("Claimant", "Client", "Jordan Lawyer") +
                 party("Mediator", "Neutral", "Jordan Lawyer"))
        r = build_party_reports([c])
        self.assertEqual(c["role"], "Multiple roles")
        self.assertEqual(len(r["clients"]["cases"]), 1)
        self.assertEqual(r["clients"]["cases"][0]["partyRoles"], ["Claimant", "Petitioner"])
        self.assertEqual(len(r["parties"]), 2)

    def test_names_do_not_fuzzily_merge_legal_entities_or_people(self):
        names = ["Uber LLC", "Uber Inc", "The Chemours Company", "Chemours Company FC LLC", "John M. Sample", "John Middle Sample", "James Davis", "Davis, James"]
        c = case(''.join(party("Defendant", n, "Jordan Lawyer") for n in names))
        self.assertEqual(len(build_party_reports([c])["clients"]["summary"]), len(names))
        self.assertEqual(party_key("  Éxample  LLC "), party_key("E\u0301XAMPLE LLC"))

    def test_client_types_are_per_case_and_opponents_are_not_clients(self):
        first = case(party("Plaintiff", "Shared Client", "Jordan Lawyer") +
                     party("Defendant", "Opponent", "Other Lawyer"))
        second = case(party("Defendant", "Shared Client", "Jordan Lawyer") +
                      party("Plaintiff", "Opponent", "Other Lawyer"), "njdc|1:24-cv-00002")
        summary, cases = report_tables(build_party_reports([first, second]))
        self.assertEqual(summary["rows"], [["Shared Client", 2, "Shared Client", ""]])
        self.assertNotIn("Client type", summary["headers"])
        rows = [dict(zip(cases["headers"], row)) for row in cases["rows"]]
        self.assertEqual({(r["Case number"], r["Client type"], r["Role"]) for r in rows}, {
            ("1:24-cv-00001", "Plaintiff", "Civil Plaintiff"),
            ("1:24-cv-00002", "Defendant", "Civil Defense")})
        self.assertEqual({r["Name"] for r in rows}, {"Shared Client"})

    def test_exports_include_rankings_cases_coverage_and_literal_source_text(self):
        c = case(party("Plaintiff", "=DANGEROUS()", "Jordan Lawyer") + party("Defendant", "Opponent", "Defense"))
        c["caseTitle"] = "<img src=x>"
        with tempfile.TemporaryDirectory() as folder:
            export_local([c], folder, "Fixture", {"generatedUtc":"2026-09-30", "lawyer":{"firstName":"Jordan","lastName":"Lawyer"}})
            book = load_workbook(Path(folder) / "case-index.xlsx")
            self.assertEqual(book.active["D6"].value, "Role")
            self.assertEqual(book["Clients"]["A4"].value, "=DANGEROUS()")
            self.assertEqual(book["Clients"]["A4"].data_type, "s")
            self.assertEqual(book["Clients"]["B4"].value, 1)
            self.assertIn("1 of 1", book["Coverage"]["A2"].value)
            self.assertEqual(set(book.sheetnames), {"Case index","Coverage","Clients","Clients cases"})
            with zipfile.ZipFile(Path(folder)/"party-reports.zip") as archive:
                self.assertEqual(len(archive.namelist()), 4)
                rows = list(csv.reader(io.StringIO(archive.read("clients-cases.csv").decode("utf-8-sig"))))
                self.assertEqual(rows[1][0], "'=DANGEROUS()")
                self.assertIn("summaries", archive.read("README.txt").decode().lower())
            html = (Path(folder)/"case-index.html").read_text(encoding="utf-8")
            self.assertNotIn("<img src=x>", html)
            self.assertIn("&lt;img src=x&gt;", html)

    def test_google_reports_use_numeric_counts_literal_names_and_coverage_notes(self):
        c = case(party("Plaintiff", "=DANGEROUS()", "Jordan Lawyer"))
        with patch('roc.output.google_token', return_value='fictional'), patch('roc.output.google_request') as request:
            request.side_effect = [{'spreadsheetId':'fictional','spreadsheetUrl':'https://example.invalid'}, {}]
            publish_google([c], "Fixture", "unused.json")
        self.assertEqual(len(request.call_args_list[0].args[1]["sheets"]), 3)
        updates = [r["updateCells"] for r in request.call_args_list[1].args[1]["requests"] if "updateCells" in r]
        clients = next(u for u in updates if u["start"]["sheetId"] == 1)
        self.assertEqual(clients["rows"][1]["values"][0]["userEnteredValue"], {"stringValue":"=DANGEROUS()"})
        self.assertEqual(clients["rows"][1]["values"][1]["userEnteredValue"], {"numberValue":1})
        self.assertIn("1 of 1", clients["rows"][0]["values"][0]["note"])
