"""Name-rule persistence, source boundaries, previews and free rebuilds."""
import contextlib
from copy import deepcopy
import io
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from roc.cli import run
from roc.exports import output_directory
from roc.common import RocError, read_json, write_json
from roc.name_rules import NameRules, normalize_rules
from roc.output import export_local
from roc.parties import build_party_reports
from roc.store import RunStore
from roc.workspace import Workspace
from tests.test_party_reports import case
from tests.test_roc import party, count


def rule(label="Example group", names=None, kind="organization-group", identifier="a"*32):
    return {"id":identifier, "label":label, "kind":kind, "names":names or ["Example Client", "Another Defendant"]}


class NameRuleTests(unittest.TestCase):
    def test_group_deduplicates_cases_but_not_counsel_or_charges(self):
        first = case(party("Defendant", "Example Client", "Jordan Lawyer", count("CLIENT COUNT", "1"), "1") +
                     party("Defendant", "Another Defendant", "Other Lawyer", count("OTHER COUNT", "2"), "2"), kind="CRIMINAL")
        second = case(party("Defendant", "Another Defendant", "Jordan Lawyer"), key="njdc|1:24-cr-00002", kind="CRIMINAL")
        original = deepcopy([first, second])
        r = build_party_reports([first,second], name_rules={"revision":1,"rules":[rule()]})
        self.assertEqual([first,second],original)
        s = r["defendants"]["summary"][0]
        self.assertEqual((s["name"], s["caseCount"], s["groupKind"]), ("Example group",2,"organization-group"))
        self.assertEqual(len(r["defendants"]["cases"]),2)
        self.assertEqual(r["clients"]["summary"][0]["caseCount"],2)
        client = next(p for p in r["clients"]["cases"] if p["caseKey"]==first["key"])
        self.assertEqual(client["sourceNames"],["Example Client"])
        self.assertEqual(client["defendantNumbers"],["1"])
        self.assertEqual(client["nature"],"CLIENT COUNT (1)")
        merged = next(p for p in r["defendants"]["cases"] if p["caseKey"]==first["key"])
        self.assertEqual({p["name"] for p in merged["sourceParties"]},{"Example Client","Another Defendant"})
        self.assertEqual(len(r["parties"]),3)

    def test_name_correction_renames_only_explicit_members(self):
        c = case(party("Plaintiff","Jane M. Sample","Jordan Lawyer") + party("Plaintiff","Jane Middle Sample","Jordan Lawyer") +
                 party("Plaintiff","Jane Sample","Jordan Lawyer"))
        rules={"revision":3,"rules":[rule("Jane Middle Sample",["Jane M. Sample"],"name-correction")]}
        r=build_party_reports([c],name_rules=rules)
        self.assertEqual({s["name"] for s in r["clients"]["summary"]},{"Jane Middle Sample","Jane Sample"})
        self.assertEqual(len(r["clients"]["cases"]),2)
        self.assertTrue(all(s["caseCount"]==1 for s in r["clients"]["summary"]))

    def test_overlaps_chains_and_invalid_rules_are_rejected(self):
        for other in (rule("Other",[" example  CLIENT "],identifier="b"*32),
                      rule("Other",["Example group"],identifier="b"*32),
                      rule("Example Client",["Different"],identifier="b"*32)):
            with self.assertRaisesRegex(RocError,"already in another rule"):
                normalize_rules([rule(),other])
        for bad in (rule(kind="guess"), rule(identifier="../file"), rule(label=""), rule(names=["x\ny"]), rule(names=[12])):
            with self.assertRaises(RocError): normalize_rules([bad])

    def test_history_survives_restart_and_undo_has_monotonic_revision(self):
        with tempfile.TemporaryDirectory() as tmp:
            store=NameRules(Path(tmp)/"name-rules.json")
            request={"revision":0,"operation":"save","rule":{k:v for k,v in rule().items() if k!='id'}}
            before,proposal=store.propose(request)
            self.assertFalse(store.path.exists())
            store.commit(request,proposal)
            store=NameRules(store.path)
            saved=store.public()
            self.assertEqual(saved["revision"],1)
            self.assertTrue(saved["canUndo"])
            edit={"revision":1,"operation":"save","rule":saved["rules"][0] | {"label":"Renamed group"}}
            store.commit(edit,store.propose(edit)[1])
            undo={"revision":2,"operation":"undo"}
            store.commit(undo,store.propose(undo)[1])
            self.assertEqual(store.public()["rules"],saved["rules"])
            self.assertEqual(store.public()["revision"],3)
            undo={"revision":3,"operation":"undo"}
            store.commit(undo,store.propose(undo)[1])
            self.assertEqual(store.public()["rules"],[])
            self.assertFalse(store.public()["canUndo"])
            with self.assertRaisesRegex(RocError,"another tab"):store.propose(request)
            store.path.write_text('{broken',encoding='utf-8')
            with self.assertRaisesRegex(RocError,"unreadable"):store.public()
            self.assertEqual(store.path.read_text(),'{broken')

    def test_exports_keep_rules_and_literal_source_associations(self):
        c=case(party("Plaintiff","Example Client","Jordan Lawyer") + party("Defendant","Another Defendant","Other Lawyer"))
        rules={"revision":2,"rules":[rule("=EXAMPLE()") ]}
        with tempfile.TemporaryDirectory() as tmp:
            export_local([c],tmp,"Fixture",{"generatedUtc":"2026-10-01","nameRules":rules})
            book=load_workbook(Path(tmp)/"case-index.xlsx")
            self.assertEqual(book["Name rules"]["A2"].data_type,"s")
            self.assertEqual(book["Clients"]["A4"].value,"=EXAMPLE()")
            self.assertIn("Example Client",book["Clients cases"]["Q4"].value)
            self.assertNotIn("Another Defendant",book["Clients cases"]["Q4"].value)
            self.assertEqual(read_json(Path(tmp)/"evidence.json")["run"]["nameRules"]["revision"],2)


class WorkspaceNameRuleTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.stdout=contextlib.redirect_stdout(io.StringIO());self.stdout.__enter__();self.addCleanup(self.stdout.__exit__,None,None,None)
        self.ws=Workspace(self.tmp.name,lambda:(_ for _ in ()).throw(AssertionError("No sign-in")))
        self.addCleanup(self.ws.close)
        self.network=patch('roc.pacer.request_json',side_effect=AssertionError('No PACER requests'));self.network.start();self.addCleanup(self.network.stop)

    def demo(self):
        identifier=self.ws.new(demo=True);self.ws.future.result(10)
        from tests.test_workspace import retrieval_values
        self.ws.act(identifier,"retrieve",retrieval_values(self.ws, identifier, ["nysdc|1:24-cr-00001"]));self.ws.future.result(10)
        return identifier

    def request(self,identifier=None):
        return {"revision":self.ws.name_rules.public()["revision"],"operation":"save","runId":identifier,
                "rule":{k:v for k,v in rule().items() if k!='id'}}

    def apply(self,request):
        preview=self.ws.preview_name_rule(request)
        result=self.ws.apply_name_rule(request | {"previewId":preview["previewId"]})
        if request.get('runId'): self.ws.future.result(10)
        return result

    def test_preview_apply_refresh_all_views_without_modifying_ledger(self):
        one=self.demo();two=self.demo()
        folder=self.ws.folder(one)
        original=read_json(output_directory(folder)/'evidence.json')['cases']
        ledger=read_json(folder/'ledger.json') if (folder/'ledger.json').exists() else None
        request=self.request(one)
        preview=self.ws.preview_name_rule(request)
        self.assertEqual(preview['affectedRuns'],2)
        clients=next(c for c in preview['comparisons'] if c['report']=='clients')
        self.assertEqual(len(clients['before']),1)
        self.assertEqual(clients['after'][0]['caseCount'],1)
        self.assertEqual(self.ws.name_rules.public()['revision'],0)
        self.assertFalse((Path(self.tmp.name)/'name-rules.json').exists())
        with self.assertRaisesRegex(RocError,'Preview'): self.ws.apply_name_rule(request)
        self.apply(request)
        self.assertEqual(read_json(output_directory(folder)/'evidence.json')['cases'],original)
        self.assertEqual(read_json(folder/'ledger.json') if (folder/'ledger.json').exists() else None,ledger)
        self.assertEqual(self.ws.summary(one,True)['partyReports']['defendants']['summary'][0]['caseCount'],1)
        self.assertFalse(self.ws.summary(one)['exportsNeedRefresh'])
        self.assertTrue(self.ws.summary(two)['exportsNeedRefresh'])
        with self.assertRaisesRegex(RocError,'Update exports'):self.ws.download(two,'party-reports.zip')
        self.ws.act(two,'refresh-reports');self.ws.future.result(10)
        self.assertFalse(self.ws.summary(two)['exportsNeedRefresh'])
        self.assertTrue(self.ws.download(two,'party-reports.zip').exists())
        newer=self.demo()
        self.assertEqual(self.ws.summary(newer,True)['partyReports']['defendants']['summary'][0]['name'],'Example group')

    def test_old_report_layout_requires_free_export_refresh(self):
        identifier=self.demo()
        folder=self.ws.folder(identifier)
        metadata=read_json(output_directory(folder)/'result.json')
        metadata.pop('clientReportVersion')
        write_json(output_directory(folder)/'result.json',metadata)
        receipts=self.ws.receipts(identifier)
        self.assertTrue(self.ws.summary(identifier)['exportsNeedRefresh'])
        with self.assertRaisesRegex(RocError,'Update exports'):
            self.ws.download(identifier,'case-index.xlsx')
        with patch('roc.pacer.request_json',side_effect=AssertionError('No PACER requests')):
            self.ws.act(identifier,'refresh-reports');self.ws.future.result(10)
        self.assertFalse(self.ws.summary(identifier)['exportsNeedRefresh'])
        self.assertEqual(self.ws.receipts(identifier),receipts)
        book=load_workbook(self.ws.download(identifier,'case-index.xlsx'))
        self.assertEqual(book.sheetnames,['Case index','Coverage','Clients','Clients cases'])

    def test_stale_preview_cannot_commit_after_case_changes(self):
        identifier=self.demo();request=self.request(identifier);preview=self.ws.preview_name_rule(request)
        path=output_directory(self.ws.folder(identifier))/'evidence.json'
        value=read_json(path);value['cases']=[];write_json(path,value)
        with self.assertRaisesRegex(RocError,'Preview'):self.ws.apply_name_rule(request | {'previewId':preview['previewId']})
        self.assertEqual(self.ws.name_rules.public()['revision'],0)

    def test_remove_undo_restore_source_names_and_old_runs(self):
        identifier=self.demo();saved=self.apply(self.request(identifier))
        removed=self.apply({'revision':saved['revision'],'operation':'delete','runId':identifier,'rule':{'id':saved['rules'][0]['id']}})
        self.assertEqual(len(self.ws.summary(identifier,True)['partyReports']['defendants']['summary']),2)
        self.apply({'revision':removed['revision'],'operation':'undo','runId':identifier})
        self.assertEqual(len(self.ws.summary(identifier,True)['partyReports']['defendants']['summary']),1)
        self.ws.close();self.ws=Workspace(self.tmp.name);self.addCleanup(self.ws.close)
        self.assertEqual(self.ws.name_rules.public()['rules'][0]['label'],'Example group')

    def test_pending_receipts_and_interrupted_work_are_not_reset_by_rules(self):
        identifier=self.demo();folder=self.ws.folder(identifier)
        with RunStore(folder,500) as store:store.reserve('pcl',{'fictional':True},10)
        m=self.ws.manifest(identifier);m.update(state='stopped',needsResume=True,lastAction='retrieve');self.ws.save(identifier,m)
        ledger=(folder/'ledger.json').read_bytes()
        self.apply(self.request(identifier))
        summary=self.ws.summary(identifier)
        self.assertEqual(summary['pendingCount'],1)
        self.assertEqual(summary['state'],'stopped')
        self.assertEqual(summary['lastAction'],'retrieve')
        self.assertEqual((folder/'ledger.json').read_bytes(),ledger)
        with self.assertRaisesRegex(RocError,'Unresolved receipts'):self.ws.quote(identifier,['nysdc|1:24-cr-00001'])

    def test_cli_uses_saved_rule_file_and_validates_before_signin(self):
        identifier=self.demo();folder=self.ws.folder(identifier)
        rules_path=Path(self.tmp.name)/'name-rules.json'
        write_json(rules_path,{'revision':1,'rules':[rule()]})
        config=read_json(folder/'config.json');config.pop('nameRules',None);config['nameRulesFile']=str(rules_path)
        write_json(folder/'config.json',config)
        self.assertEqual(run(folder/'config.json'),0)
        r=read_json(output_directory(folder)/'party-reports.json')
        self.assertEqual(r['defendants']['summary'][0]['name'],'Example group')
        write_json(rules_path,{'revision':2,'rules':[rule(),rule(identifier='b'*32)]})
        with patch('roc.pacer.Session.prompt',side_effect=AssertionError('No sign-in')):
            with self.assertRaisesRegex(RocError,'already'):run(folder/'config.json',live=True)
