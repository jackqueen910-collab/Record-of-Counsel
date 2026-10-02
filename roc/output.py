"""Present data without a chat connector: CSV, XLSX, HTML, and optional Sheets."""
import csv
from datetime import datetime
from html import escape
import json
from pathlib import Path
import urllib.parse
import urllib.request
import zipfile

from .common import RocError, write_json
from .docket import ROLE_VALUES
from .parties import CLIENT_REPORT_VERSION, build_party_reports, case_role, coverage_text, report_tables
from .review import case_issues, issue_text
from .search import counsel_aliases

HEADERS = ["Case number", "Case title", "Case Type", "Role", "Court", "District", "Date filed", "Nature of Case", "Status (PACER)", "PACER link"]
FIELDS = ["caseNumber", "caseTitle", "caseType", "role", "court", "district", "dateFiled", "nature", "status", "pacerLink"]
WIDTHS = [20, 38, 14, 21, 24, 30, 15, 40, 18, 22]


def values(case):
    return [case_role(case) if field == "role" else case.get(field, "") for field in FIELDS]


def csv_safe(value):
    value = str(value)
    return "'" + value if value.lstrip().startswith(("=", "+", "-", "@")) else value


def source_note(case):
    evidence = case.get("enrichment")
    if not evidence:
        return "Source: official PCL index. " + ("Found by: " + "; ".join(case['matchedSearchNames']) + ". " if case.get('matchedSearchNames') else '') + "; ".join(case["warnings"])
    coverage = evidence.get("courtCoverage")
    validation = (" Court retrieval validation at run: " + coverage["validationStatus"] +
                  "; live-verified samples: " + str(coverage["verifiedSamples"]) + "." if coverage else "")
    return ("Docket: " + evidence["sourceFile"] + "; SHA256: " + evidence["sourceSha256"] +
            ". Represented parties: " + ", ".join(evidence["representedParties"]) +
            ". Source roles: " + "; ".join(p["party"] + ": " + p["role"] for p in evidence.get("representedRoles", [])) +
            ". Counts copy all listed source rows, including original, superseding and terminated entries. " +
            "Defendant associations, source sections and dispositions are retained in evidence.json. " +
            "; ".join(issue_text(i) for i in case_issues(case)) + validation)


def export_local(cases, folder, title, metadata):
    litigant = metadata.get('searchType') == 'litigant'
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    with (folder / "case-index.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(HEADERS)
        writer.writerows([csv_safe(x) for x in values(c)] for c in cases)
    aliases = counsel_aliases(metadata) if metadata.get('lawyer') else []
    reports = build_party_reports(cases, aliases, metadata.get("nameRules"))
    metadata["nameRules"] = reports["nameRules"]
    metadata["partyCoverage"] = reports["coverage"]
    metadata["clientReportVersion"] = CLIENT_REPORT_VERSION
    write_json(folder / "evidence.json", {"run": metadata, "cases": cases})
    write_json(folder / "party-reports.json", reports)
    tables = ([{'name': 'Matched litigants', 'file': 'matched-litigants.csv',
                'headers': ['Name in PCL', 'PCL party role', 'Case number', 'Case title', 'Court', 'District', 'Date filed'],
                'rows': [[p['name'], p['role'], *[c.get(k, '') for k in ('caseNumber', 'caseTitle', 'court', 'district', 'dateFiled')]]
                         for c in cases for p in c.get('indexedParties', [])]}] if litigant else report_tables(reports))
    report_note = ('PCL party search: one case-index row per court/case; matched names and court-supplied role codes are preserved separately. '
                   'Names match by prefix and may refer to different people or entities. Role codes are not interpreted as proof of who filed a claim. '
                   'Docket enrichment, client reports and Document Grabber are not included for litigant searches in this first step.' if litigant else
                   coverage_text(reports) + "\n" + "\n".join(reports[k] for k in ("countPolicy", "namePolicy", "relationshipPolicy")))
    queries = metadata.get('searchQueries', [])
    if len(queries) > 1:
        report_note += '\nCombined name searches: ' + '; '.join(q['name'] for q in queries) + '. Each court/case is counted once. Grouping does not establish that names are the same person or legal entity.'
        match_rows = list(dict.fromkeys(tuple([name, ' '.join(str(row.get(k) or '') for k in ('firstName', 'middleName', 'lastName', 'generation')).strip(),
                           *[c.get(k, '') for k in ('caseNumber', 'caseTitle', 'court', 'district')]])
                           for c in cases for row in c['sourceRows'] for name in row.get('_rocSearchNames', [])))
        tables.append({'name': 'Search matches', 'file': 'search-matches.csv',
                       'headers': ['Search name', 'Name in PCL', 'Case number', 'Case title', 'Court', 'District'], 'rows': match_rows})
    for table in tables:
        with (folder / table["file"]).open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(table["headers"])
            writer.writerows([csv_safe(x) for x in row] for row in table["rows"])
    with zipfile.ZipFile(folder / "party-reports.zip", "w", zipfile.ZIP_DEFLATED) as archive:
        for table in tables:
            archive.write(folder / table["file"], table["file"])
        archive.writestr("README.txt", report_note)
        archive.writestr("name-rules.json", json.dumps(reports["nameRules"], indent=2, ensure_ascii=False))
    write_json(folder / "review.json", [{"caseNumber": c["caseNumber"], "court": c["courtId"],
               "issues": case_issues(c), "items": [issue_text(i) for i in case_issues(c)]}
               for c in cases if case_issues(c)])
    try:
        from openpyxl import Workbook
        from openpyxl.comments import Comment
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError:
        raise RocError("Install ROC's openpyxl dependency to export the workbook. CSV and evidence were saved.") from None
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Case index"
    sheet.sheet_view.showGridLines = False
    for row in [[], [title], [f"{len(cases)} cases. Generated {metadata['generatedUtc'][:10]}."],
                ["Official PCL API index; court-web docket enrichment where available. See cell notes and evidence.json."],
                [report_note if litigant else "Nature of Case copies listed counts, including original and superseding entries. Source gaps and review items are categorized in review.json."], HEADERS]:
        sheet.append(row)
    for column, width in enumerate(WIDTHS, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet["A2"].font = Font(name="Arial", size=15, bold=True)
    for row_number, case in enumerate(cases, 7):
        for col, val in enumerate(values(case), 1):
            cell = sheet.cell(row_number, col, val)
            cell.data_type = "s"  # Never interpret case text as an Excel formula.
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="center", wrap_text=col == 2)
        try:
            sheet.cell(row_number, 7, datetime.strptime(case["dateFiled"], "%Y-%m-%d"))
            sheet.cell(row_number, 7).number_format = "mm/dd/yy"
        except ValueError:
            pass
        sheet.row_dimensions[row_number].height = 28
        if case["pacerLink"].startswith("https://ecf."):
            sheet.cell(row_number, 10).hyperlink = case["pacerLink"]
            sheet.cell(row_number, 10).font = Font(name="Arial", size=10, color="1155CC")
        for col in (4, 8):
            sheet.cell(row_number, col).comment = Comment(source_note(case), "ROC")
    for cell in sheet[6]:
        cell.font = Font(name="Arial", size=10, bold=True)
        cell.fill = PatternFill("solid", fgColor="F1F3F4")
    end = max(7, len(cases) + 6)
    sheet.freeze_panes = "C7"
    sheet.auto_filter.ref = f"A6:J{end}"
    if not litigant:
        dv = DataValidation(type="list", formula1='"' + ','.join(ROLE_VALUES) + '"', allow_blank=True)
        sheet.add_data_validation(dv)
        dv.add(f"D7:D{end}")
    coverage = wb.create_sheet("Coverage")
    for line in [title, *report_note.splitlines(), *([] if litigant else ["Names are labels, not verified unique identities. Former counsel and terminated parties may be included.",
                 "Case index retains every indexed case. Client sheets include counsel-matched appearances in saved dockets only.",
                 "Client type is the listed party role for that case. A client can have different types in different cases; Role describes the lawyer's role in the case."])]:
        coverage.append([line])
        coverage.cell(coverage.max_row, 1).data_type = "s"
        coverage.cell(coverage.max_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
        coverage.row_dimensions[coverage.max_row].height = 45
    coverage.column_dimensions["A"].width = 110
    for table in tables:
        report_sheet = wb.create_sheet(table["name"])
        report_sheet.sheet_view.showGridLines = False
        report_sheet.append([report_note if litigant else coverage_text(reports)])
        report_sheet.append(["Matched source names and unmodified court role codes." if litigant else "Distinct court/case counts. See Coverage for name and relationship rules."])
        report_sheet.append(table["headers"])
        for row in table["rows"]:
            report_sheet.append(row)
        for row in report_sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.font = Font(name="Arial", size=10, bold=cell.row == 3)
                cell.alignment = Alignment(vertical="center", wrap_text=False)
                if cell.row == 3:
                    cell.fill = PatternFill("solid", fgColor="F1F3F4")
            report_sheet.row_dimensions[row[0].row].height = 25
        for col, header in enumerate(table["headers"], 1):
            report_sheet.column_dimensions[get_column_letter(col)].width = (40 if header in ("Name", "Case title", "Nature of Case", "Source names") else 26)
        report_sheet.freeze_panes = "B4"
        report_sheet.auto_filter.ref = f"A3:{get_column_letter(len(table['headers']))}{max(3, report_sheet.max_row)}"
    if reports["nameRules"]["rules"] and not litigant:
        rules_sheet = wb.create_sheet("Name rules")
        rules_sheet.append(["Preferred name", "Grouping", "Source names", "Revision"])
        for rule in reports["nameRules"]["rules"]:
            rules_sheet.append([rule["label"], rule["kind"], "; ".join(rule["names"]), reports["nameRules"]["revision"]])
        for row in rules_sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        for col in ("A", "B", "C"):
            rules_sheet.column_dimensions[col].width = 45
        rules_sheet.freeze_panes = "A2"
    wb.save(folder / "case-index.xlsx")
    headings = "".join(f"<th>{escape(h)}</th>" for h in HEADERS)
    rows = "".join("<tr>" + "".join(f'<td title="{escape(str(v), quote=True)}">{escape(str(v))}</td>' for v in values(c)) + "</tr>" for c in cases)
    navigation = '<nav><a href="#cases">Case index</a> · ' + " · ".join(f'<a href="#report-{i}">{escape(t["name"])}</a>' for i, t in enumerate(tables)) + '</nav>'
    party_html = ""
    for i, table in enumerate(tables):
        party_html += f'<h2 id="report-{i}">{escape(table["name"])}</h2><p>{escape(report_note if litigant else coverage_text(reports))}</p><table><thead><tr>' + ''.join(f'<th>{escape(h)}</th>' for h in table["headers"]) + '</tr></thead><tbody>'
        party_html += "".join('<tr>' + "".join(f'<td title="{escape(str(v), quote=True)}">{escape(str(v))}</td>' for v in row) + '</tr>' for row in table["rows"]) + '</tbody></table>'
    (folder / "case-index.html").write_text("<!doctype html><meta charset=utf-8><title>ROC reports</title><style>body{font:14px Arial;margin:28px;color:#202124}table{border-collapse:collapse;table-layout:fixed;width:2400px}th,td{border-bottom:1px solid #ddd;text-align:left;padding:9px;overflow:hidden;white-space:nowrap}th{background:#f1f3f4;position:sticky;top:0}th:nth-child(8){width:300px}h1{font-size:24px}h2{margin-top:40px}</style><h1>" + escape(title) + "</h1>" + navigation + '<p>' + escape(report_note).replace('\n', '<br>') + "</p><h2 id=cases>Case index</h2><p>" + str(len(cases)) + " cases. Hover a clipped cell to read its contents.</p><table><thead><tr>" + headings + "</tr></thead><tbody>" + rows + "</tbody></table>" + party_html, encoding="utf-8")
    return folder / "case-index.xlsx"


def google_request(url, payload, token):
    request = urllib.request.Request(url, json.dumps(payload).encode(), headers={"Content-Type": "application/json", "Authorization": "Bearer " + token}, method="POST")
    try:
        with urllib.request.urlopen(request, timeout=60) as response:
            return json.load(response)
    except Exception:
        raise RocError("Google Sheets request failed; no automatic retry. Check Drive for a partially created workbook.") from None


def google_token(credentials):
    """Use the operator's own OAuth grant; never read Codex connector credentials."""
    from .common import read_json
    creds = read_json(credentials)
    required = ("client_id", "client_secret", "refresh_token")
    if not all(creds.get(k) for k in required):
        raise RocError("Google OAuth file must contain client_id, client_secret, refresh_token.")
    body = urllib.parse.urlencode({k: creds[k] for k in required} | {"grant_type": "refresh_token"}).encode()
    try:
        with urllib.request.urlopen(urllib.request.Request("https://oauth2.googleapis.com/token", body), timeout=30) as response:
            return json.load(response)["access_token"]
    except Exception:
        raise RocError("Google OAuth refresh failed. Reauthorize the standalone application.") from None


def publish_google(cases, title, credentials, aliases=(), name_rules=None):
    # Fresh workbook per publication. Never overwrite a user's live edits.
    token = google_token(credentials)
    reports = build_party_reports(cases, aliases, name_rules)
    tables = report_tables(reports)
    sheets = [{"properties": {"sheetId": 0, "title": "Case index", "gridProperties": {"rowCount": max(1000, len(cases) + 7), "columnCount": 10, "frozenRowCount": 1, "frozenColumnCount": 2}}}]
    for i, table in enumerate(tables, 1):
        sheets.append({"properties": {"sheetId": i, "title": table["name"], "gridProperties": {"rowCount": max(1000, len(table["rows"]) + 1), "columnCount": len(table["headers"]), "frozenRowCount": 1}}})
    workbook = google_request("https://sheets.googleapis.com/v4/spreadsheets", {"properties": {"title": title}, "sheets": sheets}, token)
    identifier = workbook["spreadsheetId"]
    rows = [{"values": [{"userEnteredValue": {"stringValue": v}} for v in HEADERS]}]
    for case in cases:
        cells = [{"userEnteredValue": {"stringValue": str(v)}} for v in values(case)]
        for col in (3, 7):
            cells[col]["note"] = source_note(case)
        rows.append({"values": cells})
    requests = [{"updateCells": {"start": {"sheetId": 0, "rowIndex": 0, "columnIndex": 0}, "rows": rows, "fields": "userEnteredValue,note"}},
                {"repeatCell": {"range": {"sheetId": 0, "startRowIndex": 1, "endRowIndex": len(cases) + 1, "startColumnIndex": 7, "endColumnIndex": 8}, "cell": {"userEnteredFormat": {"wrapStrategy": "CLIP"}}, "fields": "userEnteredFormat.wrapStrategy"}},
                {"setBasicFilter": {"filter": {"range": {"sheetId": 0, "startRowIndex": 0, "endRowIndex": len(cases) + 1, "startColumnIndex": 0, "endColumnIndex": 10}}}}]
    for i, width in enumerate(WIDTHS):
        requests.append({"updateDimensionProperties": {"range": {"sheetId": 0, "dimension": "COLUMNS", "startIndex": i, "endIndex": i + 1}, "properties": {"pixelSize": width * 7}, "fields": "pixelSize"}})
    for i, table in enumerate(tables, 1):
        note = coverage_text(reports) + '\n' + '\n'.join(reports[k] for k in ("namePolicy", "relationshipPolicy"))
        report_rows = [{"values": [{"userEnteredValue": {"stringValue": h}, "note": note} for h in table["headers"]]}]
        report_rows += [{"values": [{"userEnteredValue": {"numberValue" if isinstance(v, int) else "stringValue": v}} for v in row]} for row in table["rows"]]
        area = {"sheetId": i, "startRowIndex": 0, "endRowIndex": len(report_rows), "startColumnIndex": 0, "endColumnIndex": len(table["headers"])}
        requests.extend([{"updateCells": {"start": {"sheetId": i}, "rows": report_rows, "fields": "userEnteredValue,note"}},
                         {"repeatCell": {"range": area, "cell": {"userEnteredFormat": {"wrapStrategy": "CLIP"}}, "fields": "userEnteredFormat.wrapStrategy"}},
                         {"setBasicFilter": {"filter": {"range": area}}}])
    google_request(f"https://sheets.googleapis.com/v4/spreadsheets/{identifier}:batchUpdate", {"requests": requests}, token)
    return workbook["spreadsheetUrl"]
